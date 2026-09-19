"""
extractor.py
------------
Beginner note: a "document" can arrive as a .txt, .docx, or .pdf file.
Each file type stores text differently, so this module's only job is:
    "given a file path, hand me back the plain text inside it."

Everything else in the project (validator, ML model) only ever works
with plain text, so they don't need to know or care what file format
the user originally uploaded.
"""

import os


def extract_text(file_path: str) -> str:
    """
    Reads a document from disk and returns its plain text content.

    Supported extensions: .txt, .docx, .pdf

    Raises:
        FileNotFoundError: if the path doesn't exist.
        ValueError: if the extension isn't supported.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"No such file: {file_path}")

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".txt":
        return _extract_txt(file_path)
    elif ext == ".docx":
        return _extract_docx(file_path)
    elif ext == ".pdf":
        return _extract_pdf(file_path)
    else:
        raise ValueError(
            f"Unsupported file type '{ext}'. Supported: .txt, .docx, .pdf"
        )


def _extract_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _extract_docx(file_path: str) -> str:
    # python-docx reads Word files paragraph by paragraph.
    from docx import Document

    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs]

    # Word files sometimes keep important text (like a signature block)
    # inside tables rather than normal paragraphs, so we grab those too.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.append(cell.text)

    return "\n".join(paragraphs)


def _extract_pdf(file_path: str) -> str:
    # PyPDF2 reads PDFs page by page.
    from PyPDF2 import PdfReader

    reader = PdfReader(file_path)
    pages_text = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages_text.append(text)

    return "\n".join(pages_text)
