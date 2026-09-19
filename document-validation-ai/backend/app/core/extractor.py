"""
extractor.py
------------
Extracts plain text content from documents (.txt, .docx, .pdf)
from file paths or raw bytes.
"""

import os
import io


def extract_text(file_path_or_bytes, filename: str = None) -> str:
    """
    Extracts text from a file path or raw bytes.
    Supported file extensions: .txt, .docx, .pdf
    """
    if isinstance(file_path_or_bytes, (str, bytes)) and isinstance(file_path_or_bytes, str):
        if not os.path.exists(file_path_or_bytes):
            raise FileNotFoundError(f"No such file: {file_path_or_bytes}")
        ext = os.path.splitext(file_path_or_bytes)[1].lower()
        with open(file_path_or_bytes, "rb") as f:
            content_bytes = f.read()
    else:
        content_bytes = file_path_or_bytes
        ext = os.path.splitext(filename)[1].lower() if filename else ".txt"

    if ext == ".txt":
        return content_bytes.decode("utf-8", errors="ignore")
    elif ext == ".docx":
        return _extract_docx_bytes(content_bytes)
    elif ext == ".pdf":
        return _extract_pdf_bytes(content_bytes)
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Supported formats: .txt, .docx, .pdf")


def _extract_docx_bytes(content_bytes: bytes) -> str:
    from docx import Document
    file_stream = io.BytesIO(content_bytes)
    doc = Document(file_stream)
    paragraphs = [p.text for p in doc.paragraphs]

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.append(cell.text)

    return "\n".join(paragraphs)


def _extract_pdf_bytes(content_bytes: bytes) -> str:
    from PyPDF2 import PdfReader
    file_stream = io.BytesIO(content_bytes)
    reader = PdfReader(file_stream)
    pages_text = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages_text.append(text)

    return "\n".join(pages_text)
